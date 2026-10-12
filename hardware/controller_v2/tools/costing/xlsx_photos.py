#!/usr/bin/env python3
"""Insert a 'Photo' column (C) into the Component list sheet and place one thumbnail per line.

Usage: xlsx_photos.py IN.xlsx OUT.xlsx MAPPING.json
MAPPING: {"<line number>": "<image path>", ...}; lines without an entry get a '—' in the photo cell.
Everything that refers to columns C.. is shifted one to the right: formulas on both sheets (cross-sheet
refs included), data validations, conditional formatting, column widths, the freeze pane.
"""
import json, re, sys
from pathlib import Path
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.styles import Alignment
from openpyxl.comments import Comment
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from copy import copy
from PIL import Image

SHEET = "Component list"
INSERT_AT = 3                      # new column C
THUMB_W, THUMB_H = 120, 90         # display size, px (embedded at 2x so it stays sharp when zoomed)
COL_WIDTH = 18.5                   # characters (~135 px)
ROW_HEIGHT = 74                    # points (~99 px)
EMU = 9525                         # EMU per pixel

REF = re.compile(r"(?P<sheet>(?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(?P<c1>\$?)(?P<col>[A-Z]{1,3})(?P<r1>\$?)(?P<row>\d+)")

def shift_formula(f, this_sheet):
    """Shift column refs >= INSERT_AT by one when they point at the Component list sheet."""
    def rep(m):
        sheet = m.group("sheet")
        target = sheet[:-1].strip("'") if sheet else this_sheet
        if target != SHEET: return m.group(0)
        col = column_index_from_string(m.group("col"))
        if col < INSERT_AT: return m.group(0)
        return f"{sheet or ''}{m.group('c1')}{get_column_letter(col + 1)}{m.group('r1')}{m.group('row')}"
    # don't touch string literals
    parts = re.split(r'("[^"]*")', f)
    return "".join(p if p.startswith('"') else REF.sub(rep, p) for p in parts)

def shift_range(sq):
    out = []
    for rng in str(sq).split():
        out.append(re.sub(r"\$?([A-Z]{1,3})(\$?\d+)", lambda m: (get_column_letter(column_index_from_string(m.group(1)) + 1) if column_index_from_string(m.group(1)) >= INSERT_AT else m.group(1)) + m.group(2), rng))
    return " ".join(out)

def thumb(src, dst):
    """Save a 2x-resolution copy on white; return the display size (fits THUMB_W x THUMB_H)."""
    im = Image.open(src).convert("RGBA")
    im.thumbnail((THUMB_W * 2, THUMB_H * 2), Image.LANCZOS)
    bg = Image.new("RGBA", im.size, (255, 255, 255, 255)); bg.alpha_composite(im)
    bg.convert("RGB").save(dst, "JPEG", quality=88, optimize=True)
    return max(1, round(im.width / 2)), max(1, round(im.height / 2))

def main(src, dst, mapping):
    wb = openpyxl.load_workbook(src)
    ws = wb[SHEET]
    # 1. formulas (all sheets)
    n = 0
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    new = shift_formula(c.value, sh.title)
                    if new != c.value: c.value = new; n += 1
    print("formulas shifted:", n)
    # 2. validations / conditional formatting (Component list only)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(shift_range(dv.sqref))
    cf_old = ws.conditional_formatting; cf_new = openpyxl.formatting.formatting.ConditionalFormattingList()
    for cf in cf_old:
        for rule in cf.rules:
            if rule.formula: rule.formula = [shift_formula("=" + x, SHEET)[1:] for x in rule.formula]
            cf_new.add(shift_range(cf.sqref), rule)
    ws.conditional_formatting = cf_new
    # 3. column widths
    widths = {k: v.width for k, v in ws.column_dimensions.items()}
    # 4. move cells
    ws.insert_cols(INSERT_AT)
    for k, w in widths.items():
        ci = column_index_from_string(k)
        if ci >= INSERT_AT: ws.column_dimensions[get_column_letter(ci + 1)].width = w
    ws.column_dimensions[get_column_letter(INSERT_AT)].width = COL_WIDTH
    if ws.freeze_panes:
        m = re.match(r"([A-Z]+)(\d+)", ws.freeze_panes)
        ci = column_index_from_string(m.group(1))
        if ci >= INSERT_AT: ws.freeze_panes = f"{get_column_letter(ci + 1)}{m.group(2)}"
    if ws.auto_filter and ws.auto_filter.ref: ws.auto_filter.ref = shift_range(ws.auto_filter.ref)
    if ws.print_area: ws.print_area = shift_range(ws.print_area.replace("$", ""))
    # 5. header + photos
    hdr = ws.cell(row=7, column=2); h = ws.cell(row=7, column=INSERT_AT, value="Photo")
    h.font = copy(hdr.font); h.fill = copy(hdr.fill); h.alignment = copy(hdr.alignment); h.border = copy(hdr.border)
    tmp = Path(dst).parent / "_thumbs"; tmp.mkdir(exist_ok=True)
    placed = 0
    for row in range(8, ws.max_row + 1):
        num = ws.cell(row=row, column=1).value
        if not isinstance(num, (int, float)): continue
        ln = str(int(num)); cell = ws.cell(row=row, column=INSERT_AT)
        # copy the look of the neighbouring Part cell (borders / fill)
        src_cell = ws.cell(row=row, column=2)
        cell.border = copy(src_cell.border); cell.fill = copy(src_cell.fill)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        entry = mapping.get(ln)
        if isinstance(entry, str): entry = {"file": entry}
        if entry and entry.get("file") and Path(entry["file"]).exists():
            t = tmp / f"line_{int(ln):02d}.jpg"; w, hh = thumb(entry["file"], t)
            img = XLImage(str(t)); img.width, img.height = w, hh
            col_px = int(COL_WIDTH * 7 + 5); row_px = int(ROW_HEIGHT * 96 / 72)
            marker = AnchorMarker(col=INSERT_AT - 1, colOff=max(0, (col_px - w) // 2) * EMU, row=row - 1, rowOff=max(0, (row_px - hh) // 2) * EMU)
            img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(w * EMU, hh * EMU))
            ws.add_image(img)
            ws.row_dimensions[row].height = ROW_HEIGHT
            if entry.get("comment"):
                cm = Comment(entry["comment"], "RoboCore"); cm.width, cm.height = 260, 110
                cell.comment = cm
            placed += 1
        else:
            cell.value = "—"; cell.font = copy(src_cell.font)
    print("photos placed:", placed)
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = wb.calculation or CalcProperties()
    wb.calculation.fullCalcOnLoad = True
    wb.save(dst)

if __name__ == "__main__":
    mp = json.loads(Path(sys.argv[3]).read_text()) if len(sys.argv) > 3 else {}
    main(sys.argv[1], sys.argv[2], mp)
