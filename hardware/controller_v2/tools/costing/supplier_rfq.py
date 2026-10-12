#!/usr/bin/env python3
"""Make the supplier version of the component list: a quote request for N robots (default 2).

    python3 tools/costing/supplier_rfq.py            # -> docs/RoboCore_V2_quote_request_2_robots.xlsx + .pdf

Reads quantities from docs/RoboCore_V2_component_list.xlsx (Per robot = column E, Extra = H; filament grams from
the Printed parts sheet) and photos from docs/photos (mapping.json). Leaves out the PCBWay controller, packaging,
shipping estimates, tools and any line whose quantity is 0. Internal columns (estimates, where to buy, status,
actual paid) are not copied; the specification and supplier notes are the texts in SPEC below.
Needs openpyxl and Pillow; the PDF needs Google Chrome.
"""
import argparse, json, math, subprocess, tempfile, time, datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from PIL import Image

V2 = Path(__file__).resolve().parents[2]
DOCS = V2 / "docs"
SRC = DOCS / "RoboCore_V2_component_list.xlsx"
PHOTOS = DOCS / "photos"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

EXCLUDE = {1: "controller (ordered from PCBWay)", 32: "packaging", 33: "shipping estimate", 34: "shipping estimate", 35: "tools"}

# line -> (item name, specification, usual pack, notes for the supplier, photo is only similar?)
SPEC = {
    2: ("Motor pair with encoders", "SparkFun ROB-24053 'Hobby Motor with Encoder - Plastic Gear (Pair, Red)': 2 motors, 1:48 gearbox, "
        "585 counts per wheel turn, two 6-pin cables included", "pair",
        "This exact motor, please: our chassis is built around it, so generic TT motors don't fit. Please confirm stock.", False),
    3: ("Ultrasonic distance sensor, 3.3 V", "HC-SR04P (3-5.5 V) or RCWL-1601. Board 45 x 20 mm, Ø16 mm transducers 26 mm apart, "
        "4-pin header VCC / TRIG / ECHO / GND", "1",
        "Must work from 3.3 V: the classic 5 V-only HC-SR04 is not suitable.", False),
    4: ("Line-tracking sensor, 3-channel", "3-channel TCRT5000 line-tracking module with ANALOG outputs. Pins VCC, GND, L, C, R. About 60 x 24 mm", "1",
        "Analog outputs required: digital-only modules (LM339 comparator outputs) are not suitable. Please state the board size.", True),
    5: ("Micro servo, metal gear", "MG90S, with horns and screws", "1",
        "Metal gears. Please supply a consistent, good-quality batch; we test every unit on arrival.", True),
    6: ("Micro servo", "SG90 (9 g), with horns and screws", "1", "", False),
    7: ("18650 Li-ion cell", "Genuine branded cell (Samsung, LG, Molicel, EVE or similar), about 3000 mAh, 65 mm long, unprotected", "1",
        "All cells the same model and batch. Protected cells (about 70 mm long) don't fit. Please state brand and model.", False),
    8: ("Battery holder, 3 x 18650", "3-cell 18650 holder with wire leads, about 61.6 x 78.4 mm", "1",
        "We need the cells in PARALLEL (3.7 V out). Please state whether the holder is wired in series or parallel.", True),
    9: ("1S Li-ion protection board", "1S (3.7 V) 18650 protection board, DW01 + 8205A type, rated 5 A or more", "5-10",
        "Protection only (no charger needed).", True),
    10: ("XT30 battery lead, female", "AMASS XT30U-F (female) with 10-15 cm of 16 AWG silicone wire", "5",
         "Female plug, which mates with an XT30 male (XT30UPB-M on our board).", True),
    11: ("Heat-shrink tubing assortment", "Assorted sizes, small box", "set", "", False),
    12: ("Dupont jumper wires, female-female", "2.54 mm, female-female, 20 cm", "40-wire strip", "", False),
    13: ("USB-C cable", "USB-A to USB-C or USB-C to USB-C, data + charging", "1", "Must carry data (not charge-only).", False),
    15: ("Heat-set insert, M2 (brass)", "M2 brass heat-set insert for Ø3.2 mm holes, e.g. M2 x 3 mm long, 3.2 mm OD", "50-100", "", True),
    16: ("Heat-set insert, M3 short (brass)", "M3 brass heat-set insert for Ø4.0 mm holes, e.g. M3 x 4 mm long, 4.0 mm OD", "50-100", "", False),
    17: ("Machine screw M2 x 6, pan head", "M2 x 6 mm pan head machine screw, steel", "50", "", False),
    18: ("Self-tapping screw M2 x 6, pan head", "M2 x 6 mm pan head self-tapping screw (for plastic)", "50", "", False),
    19: ("Flat washer M2", "M2 flat washer, steel", "100", "", False),
    20: ("Machine screw M2 x 8, pan head", "M2 x 8 mm pan head machine screw, steel", "50", "", False),
    21: ("Screw M2 x 5, pan head", "M2 x 5 mm pan head screw, steel", "50", "", False),
    22: ("Screw M2 x 4, pan head", "M2 x 4 mm pan head screw, steel", "50", "", False),
    23: ("Socket head cap screw M3 x 10", "M3 x 10 mm hex socket head cap screw, steel", "50", "", False),
    24: ("Self-tapping screw M3 x 8, pan head", "M3 x 8 mm pan head self-tapping screw (for plastic)", "50", "", False),
    25: ("Thumb screw M3 x 8, knurled head", "M3 x 8 mm thumb screw with knurled head", "20", "", True),
    26: ("Hex nut M8", "M8 hex nut, plain steel (about 5 g each)", "1", "", False),
    27: ("Steel ball ½\" (12.7 mm)", "12.7 mm (½\") steel bearing ball", "20", "12.7 mm glass marbles are also acceptable.", False),
    28: ("PETG filament, yellow", "PETG 3D-printer filament, 1.75 mm, yellow", "1 kg spool", "", False),
    29: ("PETG filament, black", "PETG 3D-printer filament, 1.75 mm, black", "1 kg spool", "", False),
    30: ("PETG filament, dark teal", "PETG 3D-printer filament, 1.75 mm, dark teal", "1 kg spool", "", False),
    31: ("TPU filament, black", "TPU 95A flexible filament, 1.75 mm, black", "250 g-1 kg spool", "A small spool is enough.", False),
}
SECTIONS = [("Motors, sensors and servos", range(2, 7)), ("Battery and power", range(7, 12)), ("Cables", range(12, 15)),
            ("Screws, inserts and small hardware", range(15, 28)), ("3D-printing filament", range(28, 32))]
NAVY, PALE, YEL = "1F2A44", "E8ECF3", "FFF2B3"

def read_lines(robots):
    wb = openpyxl.load_workbook(SRC)
    ws, pp = wb["Component list"], wb["Printed parts"]
    assert ws["C7"].value == "Photo" and ws["E7"].value == "Per robot", "unexpected column layout in the component list"
    waste = float(pp["C4"].value or 0)
    grams = {}                                   # (material, colour) -> (per robot, one-off)
    for r in range(7, 19):
        mat, col, per, one, mass = pp.cell(r, 5).value, pp.cell(r, 6).value, pp.cell(r, 3).value or 0, pp.cell(r, 4).value or 0, pp.cell(r, 7).value or 0
        if not mat: continue
        a, b = grams.get((mat, col), (0, 0)); grams[(mat, col)] = (a + per * mass, b + one * mass)
    fil = {28: ("PETG", "Yellow"), 29: ("PETG", "Black"), 30: ("PETG", "Dark teal"), 31: ("TPU 95A", "Black")}
    out = {}
    for r in range(8, ws.max_row + 1):
        n = ws.cell(r, 1).value
        if not isinstance(n, (int, float)): continue
        n = int(n)
        if n in EXCLUDE: continue
        if n in fil:
            per, one = grams.get(fil[n], (0, 0))
            need_g = (per * robots + one) * (1 + waste)
            out[n] = dict(qty=max(1, math.ceil(need_g / 1000)), unit="spool", note_qty=f"about {need_g:.0f} g needed")
            continue
        per = ws.cell(r, 5).value or 0; extra = ws.cell(r, 8).value or 0
        qty = per * robots + extra
        if qty <= 0: continue
        out[n] = dict(qty=int(qty), unit=ws.cell(r, 6).value or "pcs", note_qty="")
    return out

def thumb(src, dst, w=120, h=90):
    im = Image.open(src).convert("RGBA"); im.thumbnail((w * 2, h * 2), Image.LANCZOS)
    bg = Image.new("RGBA", im.size, (255, 255, 255, 255)); bg.alpha_composite(im)
    bg.convert("RGB").save(dst, "JPEG", quality=88)
    return max(1, round(im.width / 2)), max(1, round(im.height / 2))

def build_xlsx(lines, robots, out, photos):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Quote request"
    thin = Side(style="thin", color="C9D0D8"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
    widths = [5, 18.5, 26, 44, 7, 7, 13, 38, 11, 15, 11, 10, 24, 26]
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws["A1"] = "RoboCore V2 robot: request for quotation"; ws["A1"].font = F(size=16, bold=True, color=NAVY)
    ws["A2"] = f"Parts for {robots} robots. Date: {datetime.date.today().isoformat()}. Please quote in USD."; ws["A2"].font = F(italic=True, color="4A5565")
    info = [("From", "RoboCore Zambia"), ("Contact (name, email / phone)", ""), ("Deliver to", "Zambia"), ("Please reply by", "")]
    for i, (k, v) in enumerate(info):
        r = 4 + i
        ws.cell(r, 2, k).font = F(bold=True); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(r, 4, v); c.font = F(color="1F3FBF"); c.fill = PatternFill("solid", fgColor=YEL); c.border = box
    ws["H4"] = "How to fill in: the yellow columns on the right, one line per item. Partial quotes are welcome."
    ws["H5"] = "Quote the pack size you sell (e.g. a 50-pack of screws) in 'Pack / qty offered' and its price in 'Price (USD)'."
    ws["H6"] = "Photos are for identification only. * = the photo shows a similar item; the specification is what counts."
    ws["H7"] = "An M2/M3 screw, nut and insert assortment kit covering the hardware section is also welcome: please quote it separately."
    for r in range(4, 8): ws.cell(r, 8).font = F(size=9, color="4A5565")
    H = 9
    head = ["#", "Photo", "Item", "Specification", "Qty", "Unit", "Usual pack", "Notes", "Available? (Y / N / partly)",
            "Pack / qty offered", "Price (USD)", "Lead time (days)", "Your part no. / link", "Comments"]
    for i, t in enumerate(head, 1):
        c = ws.cell(H, i, t); c.font = F(bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=NAVY if i <= 8 else "7A5C00"); c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = box
    ws.row_dimensions[H].height = 32
    dv = DataValidation(type="list", formula1='"Y,N,Partly"', allow_blank=True); ws.add_data_validation(dv)
    tmp = Path(tempfile.mkdtemp()); r = H + 1; first = r; seq = 0
    for title, rng in SECTIONS:
        rows = [n for n in rng if n in lines]
        if not rows: continue
        ws.cell(r, 1, title).font = F(bold=True, color=NAVY)
        for c in range(1, 15): ws.cell(r, c).fill = PatternFill("solid", fgColor=PALE)
        r += 1
        for n in rows:
            seq += 1
            name, spec, pack, note, similar = SPEC[n]; L = lines[n]
            notes = " ".join(x for x in [(L["note_qty"].capitalize() + ".") if L["note_qty"] else "", note] if x)
            vals = [seq, None, name + (" *" if similar else ""), spec, L["qty"], L["unit"], pack, notes]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v); c.font = F(bold=(i in (3, 5))); c.border = box
                c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if i in (1, 5, 6) else "left")
            for i in range(9, 15):
                c = ws.cell(r, i); c.fill = PatternFill("solid", fgColor=YEL); c.border = box
                c.alignment = Alignment(wrap_text=True, vertical="center"); c.font = F(color="1F3FBF")
            ws.cell(r, 11).number_format = '"$"#,##0.00'
            dv.add(ws.cell(r, 9))
            p = photos.get(str(n))
            if p and Path(p).exists():
                w, h = thumb(p, tmp / f"{n}.jpg")
                img = XLImage(str(tmp / f"{n}.jpg")); img.width, img.height = w, h
                col_px, row_px = int(18.5 * 7 + 5), int(74 * 96 / 72)
                img.anchor = OneCellAnchor(_from=AnchorMarker(col=1, colOff=max(0, (col_px - w) // 2) * 9525, row=r - 1,
                                                               rowOff=max(0, (row_px - h) // 2) * 9525), ext=XDRPositiveSize2D(w * 9525, h * 9525))
                ws.add_image(img)
            ws.row_dimensions[r].height = 74
            r += 1
    last = r - 1; r += 1
    tot = [("Total for the items quoted", f"=SUM(K{first}:K{last})", False), ("Shipping to Zambia (USD)", None, True),
           ("Grand total (USD)", f"=K{r}+K{r+1}", False)]
    for i, (label, formula, fill) in enumerate(tot):
        rr = r + i
        ws.cell(rr, 8, label).font = F(bold=True); ws.cell(rr, 8).alignment = Alignment(horizontal="right")
        c = ws.cell(rr, 11, formula); c.number_format = '"$"#,##0.00'; c.font = F(bold=True, color="1F3FBF" if fill else "000000"); c.border = box
        if fill: c.fill = PatternFill("solid", fgColor=YEL)
    ws.cell(r + 4, 1, "Prices in USD. Please state any minimum order, payment terms and the shipping method.").font = F(size=9, italic=True, color="4A5565")
    ws.freeze_panes = ws.cell(H + 1, 4)
    ws.print_title_rows = f"{H}:{H}"
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    return seq

def build_pdf(lines, robots, out_pdf, photos):
    rows = []; seq = 0
    for title, rng in SECTIONS:
        ns = [n for n in rng if n in lines]
        if not ns: continue
        rows.append(f'<tr class="sec"><td colspan="7">{title}</td></tr>')
        for n in ns:
            seq += 1; name, spec, pack, note, similar = SPEC[n]; L = lines[n]
            notes = " ".join(x for x in [(L["note_qty"].capitalize() + ".") if L["note_qty"] else "", note] if x)
            p = photos.get(str(n)); im = f'<img src="file://{p}">' if p and Path(p).exists() else ""
            rows.append(f'<tr><td class="c">{seq}</td><td class="ph">{im}</td><td><b>{name}{" *" if similar else ""}</b></td><td>{spec}</td>'
                        f'<td class="c"><b>{L["qty"]}</b> {L["unit"]}</td><td class="c">{pack}</td><td>{notes}</td></tr>')
    html = f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: A4 landscape; margin: 9mm 10mm 9mm 10mm; }}
body {{ font-family: Helvetica, Arial, sans-serif; color: #1c2430; font-size: 8.6pt; }}
h1 {{ font-size: 17pt; color: #1F2A44; margin: 0 0 1mm 0; }} .sub {{ color: #4a5565; margin-bottom: 3mm; }}
.top {{ display: flex; gap: 8mm; margin-bottom: 3mm; }} .top div {{ flex: 1; }}
.kv b {{ display: inline-block; width: 42mm; color: #1F2A44; }} .kv div {{ border-bottom: .2mm solid #c9d0d8; padding: .8mm 0; }}
.how {{ background: #f3f6fa; border-left: 1mm solid #1F2A44; padding: 2mm 3mm; font-size: 8pt; line-height: 1.4; }}
table {{ border-collapse: collapse; width: 100%; }} thead {{ display: table-header-group; }}
th {{ background: #1F2A44; color: #fff; text-align: left; padding: 1.4mm; font-size: 8pt; }}
td {{ border-bottom: .2mm solid #d6dce4; padding: .45mm 1.4mm; vertical-align: middle; }}
tr {{ page-break-inside: avoid; }} td.c {{ text-align: center; white-space: nowrap; }}
td.ph {{ width: 24mm; text-align: center; }} td.ph img {{ max-width: 19mm; max-height: 11.5mm; display: block; margin: 0 auto; }}
tr.sec {{ break-after: avoid; page-break-after: avoid; }}
tr.sec td {{ background: #e8ecf3; color: #1F2A44; font-weight: bold; padding: 1mm 1.4mm; }}
.foot {{ margin-top: 3mm; font-size: 8pt; color: #4a5565; }}
</style></head><body>
<h1>RoboCore V2 robot: request for quotation</h1>
<div class="sub">Parts for {robots} robots · {datetime.date.today().isoformat()} · please quote in USD</div>
<div class="top"><div class="kv"><div><b>From</b>RoboCore Zambia</div><div><b>Contact</b>&nbsp;</div><div><b>Deliver to</b>Zambia</div><div><b>Please reply by</b>&nbsp;</div></div>
<div class="how">For each line please give: <b>available or not, the pack size you sell, its price, lead time</b> and your part number or link. Partial quotes are welcome.
Also please state the shipping cost to Zambia, any minimum order and payment terms.<br>Photos are for identification only; <b>*</b> = the photo shows a similar item, the specification is what counts.
An M2/M3 screw, nut and insert assortment kit covering the hardware section is also welcome (quote it separately). The spreadsheet version has columns to fill in.</div></div>
<table><thead><tr><th>#</th><th>Photo</th><th>Item</th><th>Specification</th><th>Qty</th><th>Usual pack</th><th>Notes</th></tr></thead><tbody>
{''.join(rows)}</tbody></table>
<div class="foot">RoboCore Zambia · quote request for {robots} robots · {seq} items</div></body></html>"""
    html_path = out_pdf.with_suffix(".html"); html_path.write_text(html)
    if out_pdf.exists(): out_pdf.unlink()
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        p = subprocess.Popen([CHROME, "--headless=new", "--disable-gpu", "--no-first-run", f"--user-data-dir={tmp}", "--allow-file-access-from-files",
                              "--no-pdf-header-footer", f"--print-to-pdf={out_pdf}", f"file://{html_path}"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        t0 = time.time(); size = -1
        while time.time() - t0 < 300:
            time.sleep(3)
            if out_pdf.exists():
                s = out_pdf.stat().st_size
                if s == size and s > 0: break
                size = s
            if p.poll() is not None and out_pdf.exists(): break
        if p.poll() is None:
            p.terminate()
            try: p.wait(timeout=10)
            except subprocess.TimeoutExpired: p.kill(); p.wait()
    html_path.unlink()
    return out_pdf.exists()

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--robots", type=int, default=2); ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()
    lines = read_lines(a.robots)
    mp = json.loads((PHOTOS / "mapping.json").read_text())
    photos = {k: str(PHOTOS / v["file"]) for k, v in mp.items()}
    out = DOCS / f"RoboCore_V2_quote_request_{a.robots}_robots.xlsx"
    n = build_xlsx(lines, a.robots, out, photos)
    print(f"{n} items -> {out}")
    skipped = sorted(set(range(1, 36)) - set(lines))
    print("left out:", ", ".join(f"{k} ({EXCLUDE.get(k, 'quantity 0')})" for k in skipped))
    if not a.no_pdf:
        ok = build_pdf(lines, a.robots, out.with_suffix(".pdf"), photos)
        print("pdf:", out.with_suffix(".pdf") if ok else "FAILED")
